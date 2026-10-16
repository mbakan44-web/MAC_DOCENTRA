import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_styled_table(doc, headers, data, header_color="1B365D"):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_color)
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=160, right=160)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)
            run.font.name = "Segoe UI"

    # Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = cell_value
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=110, bottom=110, left=160, right=160)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9.5)
                run.font.name = "Segoe UI"
                if c_idx == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(27, 54, 93)

    doc.add_paragraph() # Spacing
    return table

def add_callout(doc, text, title="STRATEJİK NOT:"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.rows[0].cells[0]
    set_cell_background(cell, "EFF5FB")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1B365D"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r1 = p.add_run(f"{title} ")
    r1.font.bold = True
    r1.font.name = "Segoe UI"
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = RGBColor(27, 54, 93)
    
    r2 = p.add_run(text)
    r2.font.italic = True
    r2.font.name = "Segoe UI"
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = RGBColor(60, 60, 60)
    doc.add_paragraph()

def build_document():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Main Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run("TÜRKİYE'NİN MERKEZİ ÜRÜN LANSMAN VE STARTUP PLATFORMU")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(27, 54, 93)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Yerli Yazılımcılar ve Girişimciler İçin 'Product Hunt Türkiye' Modeli İş Planı, Büyüme Stratejisi ve Gelir Mimarisi")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(11.5)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(100, 110, 125)

    add_callout(doc, "Bu rapor, küresel Product Hunt ve PostYourStartup platformlarının en başarılı mekaniklerini Türk yazılım ve girişimcilik ekosistemine uyarlayarak, organik viral büyüme ve kalıcı gelir elde edebilecek modern bir lansman platformunun tüm yol haritasını sunar.", "ÖZET:")

    # 1. BÖLÜM: YÖNETİCİ ÖZETİ VE PİYASA BOŞLUĞU
    h1 = doc.add_heading("1. Piyasa Analizi: Türkiye'deki Dev Fırsat ve Eksiklik", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)

    p1 = doc.add_paragraph()
    p1.add_run("Türkiye'de son yıllarda üniversite öğrencisinden kıdemli mühendisine kadar binlerce yazılımcı mobil uygulama, masaüstü yazılımı, yapay zeka aracı ve mikro-SaaS üretmektedir. Ancak bu kişilerin karşılaştığı en büyük duvar şudur: ").font.name = "Segoe UI"
    r_bold1 = p1.add_run("Ürünü bitirdikten sonra ilk 100 kullanıcıyı ve ilk geri bildirimi nereden bulacaklar?")
    r_bold1.font.bold = True
    r_bold1.font.name = "Segoe UI"

    doc.add_paragraph("Mevcut kanalların yetersizlikleri:", style='List Bullet')
    reasons = [
        ("Ekşi Sözlük: ", "Kullanıcılar yeni bir yazılım başlığı açtığında çoğu zaman linç edilmekte veya 'reklam yapma' denilerek hevesleri kırılmaktadır."),
        ("R10 ve DonanımHaber Forumları: ", "Konular saniyeler içinde aşağı kaymakta, ürün odaklı değil forum sohbeti formatında kalmaktadır."),
        ("Twitter / X & LinkedIn: ", "Geniş bir takipçi kitlesi olmayan bağımsız bir yazılımcının sesini duyurması neredeyse imkansızdır."),
        ("Yabancı Platformlar (Product Hunt vb.): ", "İngilizce bilmeyen veya hedef kitlesi yalnızca Türkiye olan (e-fatura, yerel muhasebe, Türkçe eğitim, yerli e-ticaret araçları vb.) projeler için küresel siteler sonuç vermemektedir.")
    ]
    for r_title, r_desc in reasons:
        p = doc.add_paragraph(style='List Bullet 2')
        p.add_run(r_title).font.bold = True
        p.add_run(r_desc)

    # 2. BÖLÜM: GEÇMİŞTE DENEYENLER NEDEN BAŞARISIZ OLDU?
    h2 = doc.add_heading("2. Geçmişteki Hatalar ve Bizim Çözümümüz (Neden Tutar?)", level=1)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(6)

    mistake_headers = ["Geçmişteki Hata", "Neden Başarısız Oldu?", "Bizim Kuracağımız Çözüm"]
    mistake_data = [
        [
            "Statik Liste Olmak",
            "Siteye ürün ekleyen oldu ama kimse girip bakmadı; site mezarlığa döndü.",
            "Zorunlu Oyunlaştırma (Startup Battle): Kendi ürününü eklemek isteyen herkes, önce diğer 5 yerli ürünü oylamak zorundadır."
        ],
        [
            "WhatsApp & Eş-Dost Hilesi",
            "Girişimciler akrabalarına oy bastırdı; kalitesiz ürünler 1. olunca gerçek kitle kaçtı.",
            "Ağırlıklı Puan Sistemi: Yalnızca doğrulanmış yazılımcı/girişimci profillerinin oyları sıralamada geçerli sayılır."
        ],
        [
            "Ziyaretçi Çekememek",
            "Google'da organik sıralamaya giremediler ve trafik bitti.",
            "Rozet (Backlink) Stratejisi: Ürününü tanıtan her siteye 'Yerli Girişimde Listelendi' rozeti verilir; yüzlerce siteden platformumuza devasa SEO trafiği akar."
        ]
    ]
    create_styled_table(doc, mistake_headers, mistake_data, "1B365D")

    # 3. BÖLÜM: PLATFORMUN ÇALIŞMA MEKANİZMASI
    h3 = doc.add_heading("3. Platformun 5 Aşamalı Büyüme ve Çalışma Döngüsü", level=1)
    h3.paragraph_format.space_before = Pt(14)
    h3.paragraph_format.space_after = Pt(6)

    steps = [
        ("Aşama 1 - Giriş ve Ön Eleme:", "Girişimci platforma kaydolur. Ürün adı, logo, web sitesi ve 1 cümlelik sloganını girer."),
        ("Aşama 2 - 'Girişim Savaşı' (Gamified Onboarding):", "Sistem girişimciye: 'Ürününü ücretsiz yayınlamak için önce bugün yarışan 5 yerli girişimi ikili karşılaşmalarda oyla ve 1 tanesine yapıcı yorum yaz!' kuralını uygular. Bu kural sayesinde platformda HİÇBİR YAZILIM YORUMSUZ KALMAZ; inanılmaz bir topluluk dayanışması doğar."),
        ("Aşama 3 - Günün Lansmanı & Canlı Oylama:", "Her sabah 00:01'de yeni günün ürünleri yayına girer. Kullanıcılar 'Günün Ürünü'nü seçmek için oylar. Saat 23:59'da günün ilk 3'ü belirlenir."),
        ("Aşama 4 - Doğrulanmış Rozet Sistemi (SEO Motoru):", "Girişimciye 'YerliGirişim'de Listelendi' HTML rozet kodu verilir. Girişimci bunu sitesine koyduğu an platformda 'Doğrulanmış Rozet' statüsü kazanır. Bu hamle bize binlerce ücretsiz yerli dofollow backlink kazandırır."),
        ("Aşama 5 - Haftalık Bülten ve Medya Desteği:", "Haftanın en iyi 5 yerli girişimi, teknoloji yatırımcılarına, yazılımcılara ve teknoloji basınına giden haftalık bültenimizde manşete taşınır.")
    ]
    for s_title, s_desc in steps:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(s_title).font.bold = True
        p.add_run(" " + s_desc)

    add_callout(doc, "PostYourStartup'ın dünya genelinde hızla büyümesini sağlayan en büyük sır, 'Sen başkasına değer katmadan kendi ürününü yayınlayamazsın' mantığıdır. Bu sayede platform açıldığı ilk günden itibaren canlı, hareketli ve etkileşim dolu hale gelir.", "EN KRİTİK MEKANİZMA:")

    # 4. BÖLÜM: GELİR MODELLERİ (NASIL PARA KAZANILIR?)
    h4 = doc.add_heading("4. Para Kazanma Modelleri (Monetization)", level=1)
    h4.paragraph_format.space_before = Pt(14)
    h4.paragraph_format.space_after = Pt(6)

    rev_headers = ["Gelir Kanalı", "Açıklama", "Tahmini Fiyatlandırma"]
    rev_data = [
        [
            "Günün Sponsorlu Vitrini (Featured)",
            "Ana sayfanın en tepesinde, gün boyunca 1 numaralı sabit vitrinde yer alma hakkı.",
            "Günlük 750 TL - 1.500 TL"
        ],
        [
            "Haftalık E-Posta Bülteni Sponsorluğu",
            "Yatırımcılara ve yazılımcılara giden haftalık bültenin en başında 'Haftanın Öne Çıkan Aracı' reklamı.",
            "Bülten Başına 2.500 TL - 5.000 TL"
        ],
        [
            "Doğrulanmış Dofollow Backlink Paketi",
            "Girişimcinin web sitesine kalıcı SEO katkısı sağlayan özel profil ve dizin bağlantısı.",
            "Tek Seferlik 499 TL"
        ],
        [
            "İş İlanları ve Co-Founder Köşesi",
            "Yazılımcı, tasarımcı veya ortak arayan startup'ların 30 günlük ilan listelemesi.",
            "İlan Başına 1.000 TL"
        ],
        [
            "Yatırımcı Radar Bülteni (PRO)",
            "Melek yatırımcılara 'Bu ay en hızlı büyüyen ve en çok oy alan 10 yerli yazılım' analiz raporu.",
            "Aylık Abonelik 1.500 TL"
        ]
    ]
    create_styled_table(doc, rev_headers, rev_data, "1B365D")

    # 5. BÖLÜM: MARKA VE ALAN ADI ÖNERİLERİ
    h5 = doc.add_heading("5. Marka ve İsim Önerileri", level=1)
    h5.paragraph_format.space_before = Pt(14)
    h5.paragraph_format.space_after = Pt(6)

    brands = [
        ("LansmanTR (.com / .co / .net): ", "Çok kurumsal, doğrudan ne yaptığını anlatan, akılda kalıcı ve prestijli."),
        ("UrunAvcilari (.com / .net): ", "Product Hunt'ın birebir Türkçe karşılığı, sıcak, samimi ve topluluk odaklı."),
        ("GirisimVitrini (.com): ", "Türk girişimcilerin ürünlerini sergilediği şık vitrin algısı."),
        ("StartTurk (.com / .co): ", "Uluslararası açılıma da uygun, modern ve enerjik."),
        ("YerliStartup (.com): ", "Doğrudan yerli üretime ve teknoloji milliyetçiliğine hitap eden güçlü isim.")
    ]
    for b_name, b_desc in brands:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(b_name).font.bold = True
        p.add_run(b_desc)

    # 6. BÖLÜM: 30 GÜNLÜK UYGULAMA VE YOL HARİTASI
    h6 = doc.add_heading("6. 30 Günlük Lansman ve Hayata Geçirme Yol Haritası", level=1)
    h6.paragraph_format.space_before = Pt(14)
    h6.paragraph_format.space_after = Pt(6)

    roadmap_headers = ["Hafta", "Yapılacak Görevler", "Hedef Çıktı"]
    roadmap_data = [
        [
            "1. Hafta (Mimari & Tasarım)",
            "Alan adının alınması, Next.js / Tailwind tabanlı modern karanlık mod arayüzün kurulması, veritabanı (Supabase/PostgreSQL) entegrasyonu.",
            "Çalışan prototip web sitesi"
        ],
        [
            "2. Hafta (Mekanik & Oylama)",
            "Girişim Savaşı (Startup Battle), oy verme sistemi, yorum alanı ve otomatik rozet kod oluşturucu altyapısının yazılması.",
            "Tüm oyunlaştırma özellikleri hazır"
        ],
        [
            "3. Hafta (Tohum İçerik - Seed Data)",
            "Siteyi boş açmamak için Türkiye'de bilinen ve sevilen 30 yerli girişim (Docentra, Scrile, Uzmankirala vb.) platforma önceden eklenir.",
            "Dolu ve canlı görünen bir ana sayfa"
        ],
        [
            "4. Hafta (Büyük Lansman & Yayılım)",
            "Twitter/X teknoloji topluluklarında, LinkedIn'de ve Ekşi Sözlük'te 'Türk yazılımcılar için ücretsiz lansman platformu açıldı' duyurusunun yapılması.",
            "İlk 500 aktif kullanıcı ve ilk 50 başvuru"
        ]
    ]
    create_styled_table(doc, roadmap_headers, roadmap_data, "1B365D")

    # 7. BÖLÜM: SONUÇ VE TAVSİYE
    h7 = doc.add_heading("7. Sonuç ve Stratejik Değerlendirme", level=1)
    h7.paragraph_format.space_before = Pt(14)
    h7.paragraph_format.space_after = Pt(6)

    p_final = doc.add_paragraph()
    p_final.add_run("Bu proje yalnızca bir 'web sitesi' değil; Türkiye'nin tüm yazılım ve teknoloji dünyasını tek bir çatıda toplayacak ").font.name = "Segoe UI"
    p_final.add_run("merkezi bir ekosistem ve medya gücüdür. ").bold = True
    p_final.add_run("Böyle bir platformu elinde tutan kişi, Türkiye'de hangi yeni girişimlerin doğduğunu ilk öğrenen, melek yatırımcılarla doğrudan temas kuran ve kendi diğer projelerini (Docentra gibi) hiçbir reklam ücreti ödemeden binlerce insana tek tıkla duyurabilen muazzam bir güç kazanır.")

    output_path = r"C:\Users\mustafa.bakan\Desktop\Turkiyenin_Product_Hunt_Platformu_Is_Plani_ve_Stratejisi.docx"
    doc.save(output_path)
    print(f"[SUCCESS] Document generated successfully at: {output_path}")

if __name__ == "__main__":
    build_document()
